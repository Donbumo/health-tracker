"""Read-only, bounded nutrition extraction from bytes, never domain writes.

Aliases belong here; explicit fields survive prevalidation and ambiguities
remain visible in the batch review.
PDF/image OCR runs locally with installed, bundled ONNX models. No AI provider
or external URL is called with a document, page, image or extracted diet text.
"""
import csv
import hashlib
import io
import json
import re
import stat
import threading
import unicodedata
import zipfile
from decimal import Decimal
from functools import lru_cache

from PIL import Image
from app.services.nutrition_math import NutritionError, NUTRIENTS, validate_vector
from app.services.importers.schema_detector import SchemaDetector

EXTRACTOR_VERSION = "nutrition-legacy-semantics-v3"

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED = 40 * 1024 * 1024
MAX_ROWS = 250
MAX_TEXT = 250000
OCR_LOCK = threading.Lock()


def fold(value):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(value)).lower() if not unicodedata.combining(c)).strip()


ALIASES = {
    'name': ('name','nombre','food','food name','alimento','producto','platillo','preparacion','recipe','receta'),
    'brand': ('brand','marca','brands'), 'barcode': ('barcode','codigo de barras','codigo_barras','ean','gtin'),
    'date': ('date','fecha','day','dia'), 'meal_type': ('meal_type','meal','comida','tiempo','momento'),
    'amount': ('amount','quantity','cantidad','gramos','weight_g','quantity_g','grams','peso_g'),
    'unit': ('unit','unidad','units'), 'serving_size': ('serving_size','tamano de porcion','porcion'),
    'energy': ('energy','calories','calorias','kcal','calories (kcal)','calories_kcal','calories_per_100g'),
    'protein': ('protein','proteina','proteinas','protein_g','protein (g)','protein_g_per_100g'),
    'carbohydrate_total': ('carbohydrate_total','carbohidratos totales','total carbohydrates','total carbs','total carbs (g)','total_carbs_g','carbohydrate_g','carbohydrates','carbohydrate','carbohidratos','carbs','carbs (g)','carbs_g_per_100g'),
    'net_carbs': ('net_carbs','net_carbs_g','carbohidratos netos','net_carbs_g_per_100g'),
    'fat': ('fat','grasa','grasas','fat_g','fat (g)','fat_g_per_100g'),
    'fiber': ('fiber','fibra','fiber_g','fiber (g)','fiber_g_per_100g'),
    'sugar': ('sugar','azucar','azucares','sugar_g','sugar (g)'),
    'sodium': ('sodium','sodio','sodium_mg','sodium (mg)','sodium_mg_per_100g'),
    'ingredients': ('ingredients','ingredientes','components','componentes'),
    'notes': ('notes','notas','instructions','instrucciones','preparacion_texto'),
    'basis_amount': ('basis_amount','base_cantidad'), 'basis_unit': ('basis_unit','base_unidad'),
    'equivalence': ('equivalence','equivalencia'),
    'total_mass_g': ('total_mass_g','yield_weight_g','peso_total_g','peso total g'),
    'yield_factor': ('yield_factor','factor_rendimiento','factor de rendimiento'),
}
MEALS = {'desayuno':'breakfast','breakfast':'breakfast','comida':'lunch','almuerzo':'lunch','lunch':'lunch',
         'cena':'dinner','dinner':'dinner','snack':'snack','colacion':'snack'}
for key,(label,unit,_) in NUTRIENTS.items():
    # Exact chemical forms only; generic vitamin A/folate and % are ambiguous.
    if key not in ALIASES:
        ALIASES[key]=(key, key+'_'+unit, fold(label), fold(label)+' ('+unit+')')
    elif key not in ALIASES[key]:ALIASES[key]=(key,*ALIASES[key])
    ALIASES[key]=(*ALIASES[key],key+'_'+unit+'_per_100ml',key+'_'+unit+'_per_100g')
ALIASES['energy']=(*ALIASES['energy'],'calories_per_100ml')


def bounded_rows(rows):
    rows=list(rows)
    if len(rows)>MAX_ROWS:raise NutritionError('Importación supera 250 elementos; divide el archivo.')
    return rows


def number(value):
    if isinstance(value, bool):
        raise NutritionError('Valor numérico ambiguo.')
    text = str(value).strip()
    if re.fullmatch(r'\d+,\d{3}',text):
        raise NutritionError('Separador decimal/miles ambiguo; confirma el valor.')
    text = text.replace(',','.')
    try:
        decimal = Decimal(text)
    except Exception as error:
        raise NutritionError('Valor numérico ilegible; complétalo manualmente.') from error
    if not decimal.is_finite() or decimal<0 or decimal>100000000:
        raise NutritionError('Valor fuera de rango.')
    return format(decimal,'f')


def normalized_row(raw, evidence):
    if not isinstance(raw,dict):
        raise NutritionError('Fila nutricional no estructurada.')
    folded = {fold(key):value for key,value in raw.items()}
    keys = {}
    for canonical, names in ALIASES.items():
        match = next((key for key in names if key in folded),None)
        if match is not None:
            keys[canonical]=(match,folded[match])
    row = {'evidence':evidence,'original':raw,'declaration':{},'ingredients':[],'warnings':[],
           'confirmed':False,'skip':False,'composition_complete':False}
    for key in ['name','brand','barcode','notes','serving_size','total_mass_g','yield_factor','basis_amount','basis_unit']:
        if key in keys and keys[key][1] not in (None,''):
            row[key]=str(keys[key][1])[:10000 if key=='notes' else 200]
    for key in ('date','meal_type'):
        if key in keys and keys[key][1] not in (None,''):
            value=str(keys[key][1]).strip()
            if key=='date':
                if re.fullmatch(r'\d{4}-\d{2}-\d{2}',value): row[key]=value
                else:row['planned_for']=value[:100];row['warnings'].append('Fecha o día sin calendario confirmado.')
            elif fold(value) in MEALS:
                row[key]=MEALS[fold(value)]
            else:row['meal_original']=value;row['warnings'].append('Confirma el tipo de comida.')
    if 'amount' in keys and keys['amount'][1] not in (None,''):
        raw_amount=str(keys['amount'][1]).strip()
        match=re.fullmatch(r'(\d+(?:[.,]\d+)?)\s*(g|gramos?|ml|mililitros?)',raw_amount,re.I)
        portion_match=re.fullmatch(r'(\d+(?:[.,]\d+)?)\s*(botellas?|bottles?|latas?|cans?|barras?|bars?|sobres?|sachets?|piezas?|pieces?|tortillas?|porciones?|porcion|portions?|servings?)',fold(raw_amount))
        try:row['amount']=number(match[1] if match else portion_match[1] if portion_match else raw_amount)
        except NutritionError:row['warnings'].append('Cantidad ilegible o ambigua.')
        if portion_match:row['unit']=portion_match[2]
        if match:row['unit']='ml' if fold(match[2]).startswith('m') else 'g'
        elif keys['amount'][0] in {'gramos','weight_g','quantity_g','grams','peso_g'}:row['unit']='g'
    if 'unit' in keys and keys['unit'][1] not in (None,''):
        value=fold(keys['unit'][1]);row['unit']={'gramos':'g','gramo':'g','mililitros':'ml'}.get(value,value)
    # Named portions remain portions. Only explicit equivalents can bridge to g/ml.
    portion_names = {'portion','portions','serving','servings','porcion','porciones','pieza','piezas','piece','pieces',
                     'botella','botellas','bottle','bottles','lata','latas','can','cans',
                     'barra','barras','bar','bars','sobre','sobres','sachet','sachets','tortilla','tortillas'}
    singular = {'servings':'porcion','portions':'porcion','porciones':'porcion','serving':'porcion','portion':'porcion','piezas':'pieza','pieces':'pieza','piece':'pieza',
                'botellas':'botella','bottle':'botella','bottles':'botella','latas':'lata','can':'lata','cans':'lata',
                'barras':'barra','bar':'barra','bars':'barra','sobres':'sobre','sachet':'sobre','sachets':'sobre','tortillas':'tortilla'}
    for field, label in [('unit','unit_label'),('basis_unit','basis_label')]:
        value=fold(row.get(field,''))
        if value in {'g','ml'}:row[field]=value
        if value in portion_names:
            row[label]=singular.get(value,value);row[field]='portion'
    equivalent=keys.get('equivalence',('',None))[1]
    if isinstance(equivalent,dict):
        row['equivalence']=dict(equivalent)
        row['equivalence_confirmed']=False
    from app.services.importers.nutrition_portion_text import normalize
    try:normalize(row,keys.get('amount',('',None))[1],keys.get('unit',('',None))[1],keys.get('basis_unit',('',None))[1])
    except NutritionError:row['warnings'].append('Descripción de porción ambigua: corrige cantidad, unidad o equivalencia.')
    units=raw.get('units') if isinstance(raw.get('units'),dict) else {}
    for nutrient in NUTRIENTS:
        candidate=keys.get(nutrient)
        if candidate is None:
            continue
        alias,raw_value=candidate
        if raw_value in (None,''):
            continue
        match=re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*(kcal|mg|ug|µg|μg|g)\s*',str(raw_value),re.I)
        try:value=number(match[1] if match else raw_value)
        except NutritionError:
            row['warnings'].append('Valor declarado ambiguo: '+nutrient);continue
        unit=fold(match[2]).replace('µ','u').replace('μ','u') if match else units.get(nutrient)
        if unit is not None:unit=fold(unit).replace('µ','u').replace('μ','u')
        if unit is None and (alias.endswith('_g') or alias.endswith(('_g_per_100g','_g_per_100ml')) or '(g)' in alias):unit='g'
        if unit is None and (alias.endswith('_mg') or alias.endswith(('_mg_per_100g','_mg_per_100ml')) or '(mg)' in alias):unit='mg'
        if unit is None and (alias.endswith(('_ug','_ug_per_100g','_ug_per_100ml')) or '(ug)' in alias):unit='ug'
        if unit is None and nutrient=='energy' and (alias in {'kcal','calories (kcal)','calories_kcal','calories_per_100g','calories_per_100ml'} or alias.endswith(('_kcal','_kcal_per_100ml','_kcal_per_100g'))):unit='kcal'
        row['declaration'][nutrient]={'value':value,'unit':unit,'state':'known',
            'provenance':{'method':'file_declared','source_field':alias}}
        if nutrient=='carbohydrate_total' and alias in {'carbohydrates','carbohydrate','carbohidratos','carbs','carbs (g)','carbs_g_per_100g'}:
            row['carbohydrate_proposal']=row['declaration'].pop(nutrient)
            row['warnings'].append('Carbohidratos sin distinción total/neto: clasifica explícitamente o conserva el valor sin contabilizar.')
        if unit is None:row['warnings'].append('Confirma la unidad de '+nutrient+'.')
    if isinstance(raw.get('nutrients'),dict):
        try:explicit=validate_vector(raw['nutrients'])
        except NutritionError:
            row['warnings'].append('Vector de nutrientes no reconocido: requiere un archivo estándar válido.')
            row['unsupported_vector']=True
        else:
            for nutrient,value in explicit.items():
                if nutrient in row['declaration'] and row['declaration'][nutrient]!=value:
                    row['warnings'].append('Valor declarado ambiguo: '+nutrient)
                row['declaration'][nutrient]=value
    if raw.get('composition_complete') is True:
        row['composition_complete']=True
    if 'ingredients' in keys and isinstance(keys['ingredients'][1],list):
        if len(keys['ingredients'][1])>100:raise NutritionError('Máximo 100 componentes por fila; divide la preparación.')
        row['ingredient_proposals']=[normalized_row(item,dict(evidence,ingredient=index+1)) for index,item in enumerate(keys['ingredients'][1])]
    if not row.get('name'):
        row['warnings'].append('Falta nombre; no se inventa uno.')
    per_units={unit for suffix,unit in [('per_100g','g'),('per_100ml','ml')] if any(key.endswith(suffix) for key in folded)}
    if per_units:
        per_unit=next(iter(per_units))
        explicit_conflict=False
        if row.get('basis_amount') or row.get('basis_unit'):
            try:explicit_conflict=Decimal(number(row.get('basis_amount')))!=Decimal('100') or row.get('basis_unit')!=per_unit
            except NutritionError:explicit_conflict=True
        row['basis_conflict']=len(per_units)>1 or explicit_conflict
        if not row['basis_conflict']:
            row.update(basis_amount='100',basis_unit=per_unit,kind_hint='products')
    return row


def json_rows(document, evidence):
    """Select a parent collection once; never also import its ingredient children."""
    if isinstance(document,list):return [normalized_row(row,dict(evidence,row=i+1)) for i,row in enumerate(bounded_rows(document))]
    if not isinstance(document,dict):raise NutritionError('JSON nutricional debe ser objeto o lista.')
    if document.get('schema_version')=='2.1' and document.get('type')=='nutrition_import':
        return [], document
    if document.get('record_type')=='daily_nutrition' and isinstance(document.get('data'),dict):
        data=document['data'];rows=[]
        for meal in data.get('meals',[]):
            for item in meal.get('items',[]):
                row=normalized_row(dict(item,date=data.get('date'),meal_type=meal.get('meal_type')),dict(evidence,row=len(rows)+1))
                from app.services.importers.nutrition_legacy_semantics import annotate
                annotate(row,document,meal,item)
                rows.append(row)
        if not rows:
            row=normalized_row(data,evidence);row.update(kind_hint='daily_aggregate',basis_amount='1',basis_unit='portion');rows=[row]
        return rows
    if document.get('record_type') in {'food_product','recipe'} and isinstance(document.get('data'),dict):
        return [normalized_row(document['data'],evidence)]
    for key in ('records','registros','meals','comidas','recipes','recetas','foods','alimentos','products','productos','diet','dieta','plan'):
        if isinstance(document.get(key),list):
            rows=[normalized_row(row,dict(evidence,row=i+1,section=key)) for i,row in enumerate(bounded_rows(document[key]))]
            for row in rows:
                if key in {'diet','dieta','plan'}:row['kind_hint']='plan'
                elif key in {'recipes','recetas'}:row['kind_hint']='recipes'
                elif key in {'foods','alimentos','products','productos'}:row['kind_hint']='products'
            return rows
        if isinstance(document.get(key),dict):return json_rows(document[key],dict(evidence,section=key))
    return [normalized_row(document,evidence)]


@lru_cache(maxsize=1)
def _ocr_engine():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR(det_use_cuda=False,cls_use_cuda=False,rec_use_cuda=False,
                    intra_op_num_threads=2,inter_op_num_threads=1)


def image_text(payload):
    with Image.open(io.BytesIO(payload)) as image:
        if image.width*image.height>12_000_000:
            raise NutritionError('Imagen supera 12 megapíxeles; reduce su tamaño.')
        image.load();image=image.convert('RGB')
        image.thumbnail((2200,2200))
        import numpy as np
        with OCR_LOCK:
            result,_=_ocr_engine()(np.asarray(image))
    result=result or []
    # Merge words sharing the same baseline, keeping label/value evidence.
    lines=[]
    for box,text,confidence in sorted(result,key=lambda x:(sum(p[1] for p in x[0])/4,min(p[0] for p in x[0]))):
        y=sum(p[1] for p in box)/4; height=max(p[1] for p in box)-min(p[1] for p in box)
        found=next((line for line in lines if abs(line['y']-y)<max(height*.5,8)),None)
        if found:found['parts'].append((min(p[0] for p in box),str(text)));found['confidence']=min(found['confidence'],float(confidence))
        else:lines.append({'y':y,'parts':[(min(p[0] for p in box),str(text))],'confidence':float(confidence)})
    return [{'text':' '.join(text for _,text in sorted(line['parts'])),'confidence':line['confidence']} for line in lines]


def text_rows(text,evidence):
    if len(text)>MAX_TEXT:raise NutritionError('Documento excede el límite de texto.')
    # A deliberate table/text profile: columns delimited by |, ; or tabs.
    lines=[line.strip() for line in text.splitlines() if line.strip()]
    if any('|' in line or '\t' in line for line in lines):
        delimiter='|' if any('|' in line for line in lines) else '\t'
        start=next((i for i,line in enumerate(lines) if delimiter in line and any(word in fold(line) for word in ['alimento','name','nombre','food'])),None)
        if start is not None:
            reader=csv.DictReader(io.StringIO('\n'.join(lines[start:])),delimiter=delimiter)
            return [normalized_row(row,dict(evidence,row=i+start+2)) for i,row in enumerate(bounded_rows(reader))]
    # A label is kept as one product; never treat each nutrient as another food.
    nutrients={};name=None;amount=None;unit=None;notes=[];basis_amount=None;basis_unit=None;brand=None;barcode=None;net_content=None
    aliases=[(name,key) for key,names in ALIASES.items() if key in NUTRIENTS for name in names]
    for line in lines:
        label=fold(line)
        if label.startswith(('producto:','product:','nombre:','name:')):name=line.split(':',1)[1].strip()
        if label.startswith(('marca:','brand:')):brand=line.split(':',1)[1].strip()
        if label.startswith(('barcode:','codigo de barras:')):barcode=line.split(':',1)[1].strip()
        match=re.search(r'(?:porcion|serving|cantidad|peso)\s*:?\s*(\d+(?:[.,]\d+)?)\s*(g|ml)\b',label)
        if match:amount,unit=match[1],match[2]
        match=re.search(r'(?:por|per)\s*(\d+(?:[.,]\d+)?)\s*(g|ml)\b',label)
        if match:basis_amount,basis_unit=match[1],match[2]
        if label.startswith('contenido neto'):net_content=line
        for alias,key in aliases:
            match=re.fullmatch(re.escape(alias)+r'\s*:?\s*(\d+(?:[.,]\d+)?)\s*(kcal|mg|ug|g)',label)
            if match:nutrients[alias]=match[1]+' '+match[2]
    if nutrients or name:
        raw={'name':name,'amount':amount,'unit':unit,**nutrients}
        raw['notes']='\n'.join(line for line in lines if '%' in line or fold(line).startswith(('ingredientes','ingredients','porciones por','servings per')))
        row=normalized_row(raw,evidence)
        for key,value in dict(basis_amount=basis_amount,basis_unit=basis_unit,brand=brand,barcode=barcode,net_content_text=net_content).items():
            if value:row[key]=value
        return [row]
    # Unstructured diet lines with amounts are proposals, not invented recipes.
    rows=[]
    for index,line in enumerate(lines):
        match=re.match(r'^(.+?)\s*[-:–]\s*(\d+(?:[.,]\d+)?)\s*(g|ml)\b(.*)$',line,re.I)
        if match:
            raw={'name':match[1],'amount':match[2],'unit':match[3]}
            for key,_,_ in [(k,NUTRIENTS[k][0],NUTRIENTS[k][1]) for k in NUTRIENTS]:
                aliases_for_key=ALIASES.get(key,())
                found=next(((alias,re.search(re.escape(alias)+r'\s*:?\s*(\d+(?:[.,]\d+)?)\s*(kcal|mg|ug|g)',fold(match[4]))) for alias in aliases_for_key if re.search(re.escape(alias)+r'\s*:?\s*\d',fold(match[4]))),None)
                if found and found[1]:raw[found[0]]=found[1][1]+' '+found[1][2]
            rows.append(normalized_row(raw,dict(evidence,line=index+1)))
    return rows


def extract(payload,filename,mime=None,depth=0):
    if not payload or len(payload)>MAX_BYTES:raise NutritionError('Archivo vacío o mayor de 10 MiB.')
    name=str(filename).replace('\\','/').rsplit('/',1)[-1][:200]
    evidence={'file':name,'sha256':hashlib.sha256(payload).hexdigest()}
    rows=[];texts=[];warnings=[];profile=None;standard=None
    stripped=payload.lstrip()
    if payload.startswith(b'%PDF-'):
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(payload),strict=True)
        if reader.is_encrypted or len(reader.pages)>8:raise NutritionError('PDF cifrado o mayor de ocho páginas no admitido.')
        if reader.trailer['/Root'].get('/OpenAction') or reader.trailer['/Root'].get('/AA'):
            raise NutritionError('PDF con acciones activas no admitido.')
        for index,page in enumerate(reader.pages):
            if float(page.mediabox.width)*float(page.mediabox.height)*4>12_000_000:
                raise NutritionError('Página PDF demasiado grande para OCR seguro.')
            text=page.extract_text() or '';method='pdf_text'
            if len(text.strip())<20:
                import pypdfium2 as pdfium
                pdf=pdfium.PdfDocument(payload)
                try:
                    rendered=pdf[index].render(scale=2).to_pil();buffer=io.BytesIO();rendered.save(buffer,format='PNG')
                    recognized=image_text(buffer.getvalue());text='\n'.join(item['text'] for item in recognized);method='local_ocr'
                finally:pdf.close()
            texts.append(text);rows.extend(text_rows(text,dict(evidence,page=index+1,method=method)))
        profile='pdf_text_or_local_ocr'
    elif payload.startswith((b'\x89PNG\r\n\x1a\n',b'\xff\xd8\xff')) or (payload[:4]==b'RIFF' and payload[8:12]==b'WEBP'):
        recognized=image_text(payload);text='\n'.join(item['text'] for item in recognized);texts.append(text)
        rows=text_rows(text,dict(evidence,method='local_ocr',page=1));profile='label_local_ocr'
        warnings.append('OCR local: revisa texto, unidades y cada cifra. Nada se guarda automáticamente.')
        if recognized and min(item['confidence'] for item in recognized)<.85:warnings.append('Hay texto de baja confianza; completar campos afectados.')
    elif payload.startswith(b'PK\x03\x04'):
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries=archive.infolist()
            if len(entries)>100 or sum(item.file_size for item in entries)>MAX_EXPANDED:
                raise NutritionError('ZIP supera límites de entradas/tamaño expandido.')
            for item in entries:
                if item.filename.startswith(('/', '\\')) or '..' in item.filename.replace('\\','/').split('/') or stat.S_ISLNK(item.external_attr>>16):
                    raise NutritionError('ZIP con ruta insegura o enlace.')
                if item.file_size>MAX_BYTES or (item.compress_size and item.file_size/item.compress_size>200):
                    raise NutritionError('ZIP supera límites de compresión.')
            if '[Content_Types].xml' in archive.namelist() and 'xl/workbook.xml' in archive.namelist():
                if any('vbaproject' in item.filename.lower() for item in entries):raise NutritionError('Excel con macros no admitido.')
                from openpyxl import load_workbook
                workbook=load_workbook(io.BytesIO(payload),read_only=True,data_only=False,keep_links=False)
                try:
                    if len(workbook.worksheets)>8:raise NutritionError('Excel supera ocho hojas; divide el archivo.')
                    for sheet in workbook.worksheets[:8]:
                        if sheet.max_row>MAX_ROWS+20 or sheet.max_column>60:raise NutritionError('Excel excede límites de filas/columnas.')
                        values=list(sheet.iter_rows(values_only=False))
                        start=next((i for i,row in enumerate(values) if any(fold(cell.value) in ALIASES['name'] for cell in row if cell.value is not None)),None)
                        if start is None:continue
                        header=[str(cell.value or '') for cell in values[start]]
                        for offset,cells in enumerate(values[start+1:]):
                            if any(cell.data_type=='f' for cell in cells):
                                warnings.append('Fila con fórmula no evaluada; revisar valores manualmente.');continue
                            raw={key:cell.value for key,cell in zip(header,cells) if key and cell.value is not None}
                            if raw:rows.append(normalized_row(raw,dict(evidence,sheet=sheet.title,row=start+offset+2,method='xlsx_cells')))
                    profile='xlsx_nutrition_table'
                finally:workbook.close()
            else:
                if depth:raise NutritionError('ZIP anidado no permitido.')
                profile='nutrition_export_zip'
                for item in entries:
                    if item.is_dir():continue
                    child=archive.read(item)
                    try:
                        found=extract(child,item.filename,depth=1);rows.extend(found['rows']);warnings.extend(found['warnings']);texts.append(found.get('text',''))
                    except NutritionError as error:
                        raise NutritionError('Una entrada del ZIP no es compatible; separa los archivos nutricionales antes de revisar.') from error
    elif stripped.startswith((b'{',b'[')):
        try:document=json.loads(payload)
        except (ValueError,UnicodeError) as error:raise NutritionError('JSON inválido.') from error
        # Reuse existing strict detector and assisted normalization without writes.
        if isinstance(document,dict):
            detection=SchemaDetector().detect(document)
            profile='internal_'+detection['detected_type'] if detection['mode']=='standard' else 'external_json_aliases'
        else:profile='external_json_array'
        result=json_rows(document,dict(evidence,method='json_fields'))
        if isinstance(result,tuple):rows,standard=result
        else:rows=result
    else:
        try:text=payload.decode('utf-8-sig')
        except UnicodeError as error:raise NutritionError('Firma/MIME no compatible; usa CSV UTF-8, JSON, PDF, XLSX o imagen.') from error
        if '\0' in text or len(text)>MAX_TEXT:raise NutritionError('Texto binario o demasiado largo.')
        texts.append(text)
        try:delimiter=csv.Sniffer().sniff(text[:8192],delimiters=',;\t|').delimiter
        except csv.Error:delimiter=None
        if delimiter:
            reader=csv.DictReader(io.StringIO(text),delimiter=delimiter)
            rows=[normalized_row(raw,dict(evidence,row=i+2,method='delimited_cells')) for i,raw in enumerate(bounded_rows(reader))]
            profile='myfitnesspal_csv' if {'Date','Meal','Food Name','Calories'}<=set(reader.fieldnames or []) else 'nutrition_delimited_table'
        else:rows=text_rows(text,dict(evidence,method='plain_text'));profile='nutrition_text'
    if len(rows)>MAX_ROWS:raise NutritionError('Importación supera 250 elementos; divide el archivo.')
    text='\n'.join(texts)
    explicit_destination = document.get('destination') if profile in {'external_json_aliases','external_json_array'} and isinstance(document,dict) else None
    plan_words={'dieta','prescrito','plan alimentario','meal plan','prescribed'}
    hints={row.get('kind_hint') for row in rows if row.get('kind_hint')}
    hint='plan' if any(word in fold(text) for word in plan_words) else 'products' if profile=='label_local_ocr' else next(iter(hints)) if len(hints)==1 else 'consumed' if profile=='myfitnesspal_csv' else 'ambiguous'
    if explicit_destination in {'consumed','plan','products','recipes','daily_aggregate'}:hint=explicit_destination
    if not rows and standard is None:warnings.append('No se reconocieron filas con este perfil. Requiere estructurar el texto o completar manualmente; no se inventaron alimentos.')
    return {'rows':rows,'standard':standard,'profile':profile,'classification_hint':hint,'warnings':warnings,
            'text':text[:MAX_TEXT],'source':evidence,'mime_claim':str(mime or ''),'signature_detected':True}
